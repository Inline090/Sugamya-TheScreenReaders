"""Build the final evidence deck from real local observations; no invented customer data."""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN
from pptx.oxml import parse_xml
from pathlib import Path
import json
import re
from PIL import Image
R=Presentation(); R.slide_width=Inches(13.333);R.slide_height=Inches(7.5)
INK='17242C'; BLUE='1D5FBF'; BG='F2F4F7'; MUTED='53636E'; WHITE='FFFFFF'; GREEN='235D39'; RED='9D2731'; BORDER='D5DDE3'
m=json.loads(Path('evidence/measurements.json').read_text());proof=json.loads(Path('evidence/public-deployment-proof.json').read_text())
sast_text=Path('evidence/static-analysis.txt').read_text()
sast_match=re.search(r'0 errors, (\d+) warnings',sast_text)
sast_warnings=int(sast_match.group(1)) if sast_match else 'see report'
repo='https://github.com/shivanshjain456/Sugamya-TheScreenReaders'
def box(s,x,y,w,h,fill=WHITE,line=None,radius=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=RGBColor.from_string(fill);sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=RGBColor.from_string(line)
 sh._element.spPr.append(parse_xml('<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>'))
 return sh
def text(s,x,y,w,h,txt,size=20,color=INK,bold=False,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.word_wrap=True;tf.margin_left=0;tf.margin_right=0;tf.margin_top=0;tf.margin_bottom=0
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.font.name='Arial';p.font.size=Pt(size);p.font.color.rgb=RGBColor.from_string(color);p.font.bold=bold;p.space_after=Pt(10)
  if align is not None:p.alignment=align
 return sh
def new(k,title,sub=''):
 s=R.slides.add_slide(R.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=RGBColor.from_string(BG);text(s,.55,.35,12.2,.3,'Sugamya / '+k,12,MUTED);text(s,.55,.9,12.2,.8,title,30,INK,True)
 if sub:text(s,.55,1.72,12.15,.6,sub,16,MUTED)
 box(s,.55,7.02,12.2,.012,BORDER);text(s,.55,7.15,11,.2,'The Screen Readers · HX-P3JQ · AgentOps’26 · IT Infrastructure',10,MUTED);text(s,12,7.12,.7,.23,str(len(R.slides)),11,MUTED,align=PP_ALIGN.RIGHT);return s
def note(s,txt):s.notes_slide.notes_text_frame.text=txt
def metric(s,x,y,w,value,label,caption=''):
 text(s,x,y,w,.7,value,38,BLUE);text(s,x,y+.85,w,.5,label,17,INK,True)
 if caption:text(s,x,y+1.4,w,.8,caption,14,MUTED)
def link(s,x,y,w,label,url):
 sh=text(s,x,y,w,.4,label,11,MUTED);sh.text_frame.paragraphs[0].runs[0].hyperlink.address=url
# 1
s=new('01 / Operator outcome','Verified repair is not finished until visitors receive it.','Sugamya: A Self-Verifying AI Agent That Fixes Web Accessibility')
text(s,.6,2.55,12.0,1.2,'An owner-controlled accessibility repair control plane.\nNo target source patch. No visitor extension.',25,INK,True)
metric(s,.6,4.35,3.4,'4 → 0','Seeded axe instances','One owned public fixture; not conformance')
metric(s,4.7,4.35,3.4,'0/3 → 3/3','Automated service-task trials','Same accessible-name + keyboard method')
metric(s,8.8,4.35,3.4,'11 / 11','Public visitor gates','Fresh browser, anonymous HTTPS')
note(s,'0:00–0:25. We are building for the people who operate sites, not asking disabled citizens to install another tool. Our narrow proof is a reviewed response-level repair that reaches a clean public browser. All numbers are our owned fixtures, not customer studies.')
#2
s=new('02 / Buyer and bottleneck','The bottleneck is safe delivery, not another list of issues.','A scanner or a patch on a copy does not change what an ordinary citizen receives.')
roles=[('Purchaser','Service-owning agency or managed-service provider'),('Operator','Small web operations / infrastructure team'),('Approver','Accessibility or content owner; decides meaning'),('Visitor','Uses the service normally; installs nothing')]
for i,(a,b)in enumerate(roles):
 y=2.6+i*.8;text(s,.65,y,2.2,.4,a,19,BLUE,True);text(s,3.05,y,9.2,.55,b,20)
text(s,.65,6.1,11.7,.45,'Required authority: control the origin installation and visitor routing—not access to application source.',17,INK,True)
note(s,'0:25–0:50. We preserve the submitted problem and target operator. Infrastructure staff must control a hostname/routing layer or obtain an owner-installed integration. Without that authority Sugamya cannot modify an arbitrary website. This is why IT Infrastructure is the selected current portal domain.')
#3
s=new('03 / Landscape','We are not claiming the first accessibility agent.','The differentiated proof is the bounded change-control loop, not AI, PRs, alt text, or runtime repair alone.')
rows=[('GitHub agent / Scanner','Objective code remediation, issues, Copilot PR workflow'),('GitHub alt-text plugin','Deterministic quality rules + opt-in image/context model review'),('AudioEye','Site-installed runtime fixes plus expert/custom remediation'),('axe / enterprise platforms','Strong automated evidence, guided tests and governance')]
for i,(a,b)in enumerate(rows):
 y=2.6+i*.65;text(s,.65,y,3.1,.45,a,17,INK,True);text(s,3.9,y,8.65,.5,b,17,MUTED)
box(s,.65,5.5,12,.85,WHITE);text(s,.85,5.7,11.6,.5,'Our claim: inspectable owner-controlled manifests + human semantics + independent visitor gates + replayable failures.',18,BLUE,True)
link(s,.65,6.52,4.3,'GitHub agent + scanner: current primary sources','https://github.blog/ai-and-ml/github-copilot/building-a-general-purpose-accessibility-agent-and-what-we-learned-in-the-process/')
link(s,5.05,6.52,3.7,'AudioEye: current runtime description','https://www.audioeye.com/post/how-audioeyes-accessibility-technology-works/')
link(s,9,6.52,3.6,'Full cited matrix in repository',repo+'/blob/main/docs/COMPETITIVE-EVIDENCE.md')
note(s,'0:50–1:20. GitHub already offers strong contextual code-based remediation and alt-text work; AudioEye already reaches ordinary visitors through runtime changes. We read the official scanner/plugin code and do not assert missing competitor features where unverified. Our deliverable is a narrow OSS operational delivery pattern with an inspectable proof. No superiority benchmark against these vendors was run.')
#4
s=new('04 / Delivery architecture','Owner-installed proxy. Approved manifest. Real visitors.','The repair engine consumes HTTP evidence; the original application remains unchanged.')
for x,w,title,body in [(0.65,3.25,'Separate origin','Owned public HTML\nExisting CSP + form endpoint'),(4.7,3.25,'Sugamya gateway','Exact page fingerprint\nAllowlisted attributes/class'),(8.75,3.85,'Ordinary browser','No extension or operator state\nVersion header + completed task')]:
 box(s,x,2.75,w,1.8,WHITE,line=BORDER);text(s,x+.18,2.96,w-.36,.45,title,20,INK,True);text(s,x+.18,3.57,w-.36,.8,body,16,MUTED)
for x in [4.1,8.15]:text(s,x,3.22,.5,.5,'→',27,BLUE,True)
text(s,.7,4.95,11.9,.6,'Observe → prioritize → semantic review → canary → activate → fresh-browser verify → keep / rollback',21,BLUE,True)
text(s,.7,5.85,11.9,.75,'Temporary Quick Tunnel proves public HTTPS reach; not production hosting.\nContrast uses an existing owner-approved utility class; dynamic/private pages are outside the released repair scope.',15,MUTED)
note(s,'1:20–1:50. A reverse proxy was selected because no domain/Worker credentials were available. It does not inject generated JavaScript. ID-only operations check exact content and attributes; mismatches preserve original HTML. A Quick Tunnel makes the actual route publicly accessible for this demonstration. Contrast needs a suitable existing class; we disclose that precondition, not claim universal styling repairs.')
#5
s=new('05 / Criterion choice','55 criteria inventoried does not mean 55 automated checks.','WCAG 2.2 AA · applicability and test methods are disclosed, not converted into a compliance score.')
vals=[('1','Selected automated','Rendered contrast evidence'),('3','Human-assisted','Alt, labels, name/role/value'),('3','Manual remainder','Full keyboard/focus assessment'),('5','Inapplicable here','No media in this static state'),('43','Unresolved / untested','No criterion-wide assessment')]
for i,(n,label,desc)in enumerate(vals):
 x=.65+i*2.5;metric(s,x,2.7,2.25,n,label,desc)
text(s,.7,5.35,11.85,.55,'Priority favors the blocked application task—not easy-to-count rules.',22,INK,True)
text(s,.7,6.02,11.85,.6,'Severity × confidence × evaluated-task reach ÷ estimated review effort. Reach and effort are disclosed assumptions.\nNo universal “20 manual criteria” partition was verified. No whole-criterion or site conformance claim.',14,MUTED)
link(s,.7,6.64,11.6,'W3C WCAG 2.2 and WAI evaluation-tool limits','https://www.w3.org/WAI/test-evaluate/tools/selecting/')
note(s,'1:50–2:15. The W3C-derived 55 is WCAG 2.2 A/AA. Tool coverage is method dependent; we did not inherit an unverified fixed 20 manual number. The image description is lower confidence than the task-blocking required field and submit action. The unsupported/unresolved list remains a manual-review queue, not a pass.')
#6
s=new('06 / Responsible AI','The model drafts text. A human decides meaning.','One real Groq request used only redacted nearby text; no image or screenshot went to a model.')
text(s,.65,2.6,5.8,.5,'Actual AI draft',20,BLUE,True);text(s,.65,3.18,5.8,1,'“Diagram showing route from entrance to service desk.”',24)
text(s,.65,4.46,5.8,.8,'Uncertainty: “Based on provided context; no visual details confirmed.”',17,MUTED)
box(s,6.95,2.5,5.7,3.3,WHITE,line=BORDER);text(s,7.15,2.7,5.3,.45,'Human edited and explicitly approved',20,INK,True);text(s,7.15,3.33,5.3,1.2,'“From Entrance, follow the path right, turn upward, then right to Service desk.”',23);text(s,7.15,4.85,5.3,.65,'All 3 semantic operations bind approval to the exact manifest; edits clear prior approvals.',15,MUTED)
text(s,.65,6.05,12,.65,'441 tokens · 380 ms observed · billed cost unverified. Provider retention exceptions may apply; ZDR not verified.\nSigned-in routes denied; no screenshot, image, full HTML, form value or credential in the model payload.',14,MUTED)
note(s,'2:15–2:45. Optional remote processing requires a preview and per-request payload consent. The real model output is retained. The owner explicitly confirmed the three exact meanings through Arena; engineering transmitted that decision through the approver role. The independent task checker cannot establish image meaning; an automated pass never overrides unresolved semantics. Groq retention exceptions are disclosed.')
#7
s=new('07 / Fresh-browser proof','The improved task reaches a public, clean browser.','Recorded evidence: actual HTTPS proxy response, accessible-name snapshot, keyboard submission and unchanged origin.')
metric(s,.65,2.6,3.1,'4 → 0','Seeded axe instances','1 page, 4 instances; no conformance inference')
metric(s,.65,4.75,3.1,'0/3 → 3/3','Same automated task','Fresh context per trial; not a screen-reader user study')
s.shapes.add_picture('evidence/public-visitor-after.png',Inches(4.25),Inches(2.43),width=Inches(8.3),height=Inches(4.32))
link(s,.65,6.55,12,'Temporary public proof URL (may expire) · full response/task evidence in repository',proof['publicURL'])
note(s,'2:45–3:15. We observed four seeded axe violation instances before and zero after, then used the same accessible-name and keyboard task method for three trials per state. All before attempts blocked; all after attempts completed. Public proof checks the actual HTTPS route from a new browser with no session or extension. This is engineering task evidence, not measured disabled-user experience.')
#8
s=new('08 / Failure and security','Passing a scan is not permission to promote.','A deliberately wrong approved test label passed syntax/axe checks—but failed the independent task gate.')
text(s,.65,2.55,5.4,.6,'Controlled negative canary',21,RED,True);text(s,.65,3.25,5.3,1.15,'“Transfer funds” instead of\n“Submit application” → withheld.',23)
text(s,.65,4.65,5.3,1.25,'Separate stale-origin test: fingerprint mismatch → original response → failed independent check → actual active-pointer rollback.',18,MUTED)
s.shapes.add_picture('evidence/negative-canary-frame.png',Inches(6.4),Inches(2.5),width=Inches(6.2),height=Inches(3.95))
text(s,.65,6.33,11.9,.5,f'30 unit tests · 21 E2E checks · 0 known dependency advisories · {sast_warnings} static warnings retained and reviewed.',14,MUTED)
note(s,'3:15–3:45. The video negative case uses an explicitly labeled engineering fixture, not fabricated human approval. Actual non-promotion is measured by unchanged visitor response. A separate stale-page test causes real rollback and emergency disable restores original HTML. Security release is narrow: fixed authorized origins, role-gated decisions, no generated code, scoped evidence, bounded requests, hash-chained history; no perfect-security or multi-tenant production claim.')
#9
s=new('09 / AI to ROI','Measure the workflow. Label the forecast.','Observed execution evidence and assumed business savings are kept separate.')
metric(s,.65,2.65,3.7,f"{m['runtime']['increment']['medianMs']:.2f} ms",'Median added local latency','30 paired warm requests; proxy + parser + local network')
metric(s,4.85,2.65,3.65,'₹3,000 / mo','Illustrative effort value','10 sites × (45−15 min) × ₹600/hour; assumptions')
metric(s,9.05,2.65,3.55,'₹300 / ₹1,500','Sensitivity cases','1 site / doubled review burden; hosting and audits excluded')
box(s,.65,5.25,12,.92,WHITE);text(s,.85,5.44,11.6,.6,'Human review/setup time was not timed. No measured saving or customer ROI is claimed.\nSecond installed target + separate origin/gateway/audit store passed the same 21 checks.',16,INK)
text(s,.7,6.48,11.9,.35,'Main model call: 441 tokens; post-hoc public-list-price estimate ≈ US$0.0000668, not an invoice.',13,MUTED)
note(s,'3:45–4:15. The measured median overhead is local paired request latency, not a geographic production benchmark. The rupee calculation is a transparent forecast with explicit times/rate and sensitivity. No timed operator study was conducted, so we do not claim measured time savings. Reuse is real on a second configuration of the same form shape, not a claim that all sites adapt automatically.')
#10
s=new('10 / Reusable CGI delivery pattern','A narrow accelerator judges can inspect and rerun.','One repeatable contract: owner authority + approved change + visitor proof + operational reversal.')
items=[('Reuse','Configuration, strict contracts and a second independent installation'),('Governance','Explicit semantic consent/approval, minimal evidence, roles and trace'),('Business value','Task outcomes and measured latency; forecast-only effort model'),('MVP readiness','Actual public route, tested controls, MP4, reproducible repository'),('Innovation','The inspectable delivery/verification combination—not AI or overlays alone')]
for i,(a,b)in enumerate(items):
 y=2.5+i*.62;text(s,.65,y,2.4,.4,a,17,BLUE,True);text(s,3.1,y,9.3,.5,b,17)
link(s,.65,6.16,12,repo,repo)
text(s,.65,6.62,11.9,.3,'Limits: fixture only · static public HTML · existing contrast class · no tenant isolation/SSO · temporary tunnel · no conformance claim',11,MUTED)
note(s,'4:15–5:00. Close on the central answer: source sharing is unnecessary if an authorized owner can route visitor responses through a tightly controlled layer. If no owner control exists, this system cannot modify the website. Ask judges to inspect the actual manifest, public response, human decision, failed canary and hash-chain replay. Transition to the five-minute live demo; recording is explicitly labeled fallback evidence, never passed off as live.')
# appendices
s=new('Appendix / Tough questions','What we will not claim.','Five-minute Q&A preparation; not part of the five-minute slide narrative.')
qas=[('“Is this source repair?”','No. It transforms eligible HTTP responses. Origin source stays unchanged; removing the route removes repairs.'),('“Why not GitHub?”','Use GitHub when the repo-owning team can merge and deploy. Our route addresses a different permission bottleneck; no superiority claim.'),('“Can it understand every image?”','No image is sent in our model path. Human review supplies visual meaning; insufficient context stays unresolved.'),('“What about all 55 criteria?”','A method-disclosing inventory is not coverage. 43 untested and 3 fuller-manual criteria remain open.'),('“Production-ready for any site?”','No. Owned static fixture release; stronger isolation, client routing, new task policies and manual audits remain prerequisites.')]
for i,(a,b)in enumerate(qas):
 y=2.35+i*.85;text(s,.65,y,3.35,.62,a,16,BLUE,True);text(s,4.25,y,8.1,.72,b,16)
note(s,'Answer directly and point to evidence/limitations. Do not extrapolate fixture numbers or competitor marketing claims into customer outcomes. Do not promise to fix signed-in or arbitrary SPA states.')
s=new('Appendix / Evidence map','Evidence, commands and fallback.','Use the repository artifacts to replay observations; use actual live controls only when the environment is running.')
for i,line in enumerate(['evidence/public-deployment-proof.json — ordinary HTTPS browser, version and task','evidence/human-approved-deployment.json — explicit project-owner semantic decision','evidence/reuse/e2e-results.json — 21-case independent reuse/failure suite','evidence/model-live-result.json — actual consented request/output and held deployment','evidence/measurements.json — denominators, 30 latency pairs and 3+3 task trials','npm test · npm run test:e2e · npm run test:ui · npm run replay','artifacts/Sugamya-demo.mp4 — actual 79.6-second captioned recording, not live']):
 text(s,.65,2.5+i*.52,12,.45,line,16)
link(s,.65,6.48,12,'Setup / architecture / security / coverage / competitor primary sources',repo+'/blob/main/README.md')
note(s,'Fallback: say “This is a recording from the tested implementation” and play the MP4. Never narrate a cached action as a live deployment. The receipt and final commit SHA are supplied separately after portal save, avoiding a circular commit.')
Path('artifacts').mkdir(exist_ok=True);R.save('artifacts/Sugamya-final.pptx')
# Structural validation, not a false claim of desktop PowerPoint rendering.
for slide in R.slides:
 for shape in slide.shapes:
  assert shape.left>=0 and shape.top>=0 and shape.left+shape.width<=R.slide_width+Inches(.01) and shape.top+shape.height<=R.slide_height+Inches(.01)
assert Path('artifacts/Sugamya-final.pptx').stat().st_size<25_000_000
print('Built',len(R.slides),'slides;',Path('artifacts/Sugamya-final.pptx').stat().st_size,'bytes. First 10 slides = five-minute narrative; 2 Q&A appendices.')
